from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .config import project_path
from .db import AtlasDB

EXPORTS = {
    "cancer_types": "SELECT * FROM cancer_types ORDER BY main_type, level, name",
    "drug_products": "SELECT * EXCLUDE(raw_json) FROM drug_products ORDER BY active_ingredient, brand_name",
    "drug_approvals": "SELECT * FROM drug_approvals ORDER BY action_date DESC NULLS LAST",
    "drug_labels": "SELECT * EXCLUDE(raw_json) FROM drug_labels ORDER BY effective_time DESC NULLS LAST",
    "clinical_trials": "SELECT * EXCLUDE(raw_json) FROM clinical_trials ORDER BY nct_id",
    "biomarker_therapy": "SELECT * EXCLUDE(raw_json) FROM biomarker_therapy_evidence ORDER BY gene, alteration, oncotree_code",
    "guidelines": "SELECT * FROM guideline_recommendations ORDER BY source, oncotree_code, stage, line_of_therapy",
    "regimens": "SELECT * FROM regimen_components ORDER BY source, protocol_id, regimen_name, component_drug",
    "licensed_monographs": "SELECT * FROM licensed_monographs ORDER BY source, drug_name",
    "coverage": "SELECT * FROM atlas_cancer_coverage ORDER BY main_type, level, name",
    "mapping_issues": "SELECT * FROM mapping_issues ORDER BY status, entity_type, source_text",
}


def _autosize(ws, max_width: int = 60) -> None:
    for col_idx, column in enumerate(ws.columns, 1):
        width = min(max((len(str(cell.value)) if cell.value is not None else 0) for cell in column) + 2, max_width)
        ws.column_dimensions[get_column_letter(col_idx)].width = max(10, width)


def run(cfg: dict[str, Any], db: AtlasDB) -> Path:
    out_dir = project_path(cfg, cfg["project"]["export_dir"])
    out_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, sql in EXPORTS.items():
        df = db.dataframe(sql)
        df.to_csv(out_dir / f"{sheet_name}.csv", index=False)
        try:
            df.to_parquet(out_dir / f"{sheet_name}.parquet", index=False)
        except Exception:
            pass
        ws = wb.create_sheet(sheet_name[:31])
        ws.append(list(df.columns))
        for row in df.itertuples(index=False, name=None):
            ws.append([str(v) if isinstance(v, (dict, list)) else v for v in row])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        _autosize(ws)
    path = out_dir / "all_cancer_treatment_atlas.xlsx"
    wb.save(path)
    return path
